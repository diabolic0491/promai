import json
import re
from collections.abc import Iterator, Mapping
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.parts.hdrftr import FooterPart, HeaderPart
from docx.text.paragraph import Paragraph


PLACEHOLDER_PATTERN = re.compile(
    r"{{\s*([A-Za-z0-9_.-]+)\s*}}"
)

RUSSIAN_MONTHS = (
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)


class MissingTemplateVariablesError(Exception):
    def __init__(
        self,
        variable_names: list[str],
    ) -> None:
        self.variable_names = variable_names
        super().__init__(
            ", ".join(variable_names)
        )


class InvalidDocxTemplateError(Exception):
    """DOCX-шаблон нельзя прочитать или сохранить."""


def flatten_form_data(
    form_data: Mapping[str, Any],
    *,
    prefix: str = "",
) -> dict[str, Any]:
    flattened: dict[str, Any] = {}

    for raw_key, value in form_data.items():
        key = str(raw_key).strip()

        if not key:
            continue

        variable_name = (
            f"{prefix}.{key}"
            if prefix
            else key
        )

        if isinstance(value, Mapping):
            flattened.update(
                flatten_form_data(
                    value,
                    prefix=variable_name,
                )
            )
        else:
            flattened[variable_name] = value

    return flattened


def format_date(value: date | None) -> str | None:
    if value is None:
        return None

    return value.strftime("%d.%m.%Y")


def add_approval_date_variables(
    values: dict[str, Any],
    approval_date: date | None,
) -> None:
    if approval_date is None:
        return

    values.update(
        {
            "approval.date": format_date(
                approval_date
            ),
            "approval.day": (
                f"{approval_date.day:02d}"
            ),
            "approval.month": RUSSIAN_MONTHS[
                approval_date.month - 1
            ],
            "approval.year": str(
                approval_date.year
            ),
        }
    )


def value_to_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")

    if isinstance(value, date):
        return format_date(value) or ""

    if isinstance(value, bool):
        return "Да" if value else "Нет"

    if isinstance(value, Decimal):
        return format(value, "f")

    if isinstance(value, (list, tuple, set)):
        return "\n".join(
            value_to_text(item)
            for item in value
        )

    if isinstance(value, Mapping):
        return json.dumps(
            value,
            ensure_ascii=False,
        )

    return str(value)


def is_missing_value(value: Any) -> bool:
    if value is None:
        return True

    if isinstance(value, str):
        return not value.strip()

    if isinstance(value, (list, tuple, set)):
        return not value

    return False


def normalize_required_variable(
    raw_variable: str,
) -> str:
    normalized = raw_variable.strip()
    match = PLACEHOLDER_PATTERN.fullmatch(
        normalized
    )

    if match is not None:
        return match.group(1)

    return normalized


def iter_document_paragraphs(
    document: DocxDocument,
) -> Iterator[Paragraph]:
    roots: list[Any] = [document._element]

    roots.extend(
        part.element
        for part in document.part.package.parts
        if isinstance(
            part,
            (HeaderPart, FooterPart),
        )
    )

    seen_roots: set[int] = set()
    seen_paragraphs: set[int] = set()

    for root in roots:
        root_id = id(root)

        if root_id in seen_roots:
            continue

        seen_roots.add(root_id)

        for paragraph_element in root.xpath(
            ".//w:p"
        ):
            paragraph_id = id(paragraph_element)

            if paragraph_id in seen_paragraphs:
                continue

            seen_paragraphs.add(paragraph_id)
            yield Paragraph(
                paragraph_element,
                root,
            )


def get_template_variables(
    document: DocxDocument,
) -> set[str]:
    variables: set[str] = set()

    for paragraph in iter_document_paragraphs(
        document
    ):
        variables.update(
            match.group(1)
            for match in PLACEHOLDER_PATTERN.finditer(
                paragraph.text
            )
        )

    return variables


def replace_paragraph_placeholders(
    paragraph: Paragraph,
    values: Mapping[str, str],
) -> None:
    runs = paragraph.runs

    if not runs:
        return

    full_text = "".join(
        run.text
        for run in runs
    )
    matches = list(
        PLACEHOLDER_PATTERN.finditer(full_text)
    )

    for match in reversed(matches):
        variable_name = match.group(1)
        replacement = values[variable_name]
        start_position = match.start()
        end_position = match.end()

        current_position = 0
        start_run_index: int | None = None
        end_run_index: int | None = None
        start_offset = 0
        end_offset = 0

        for run_index, run in enumerate(runs):
            next_position = (
                current_position + len(run.text)
            )

            if (
                start_run_index is None
                and start_position < next_position
            ):
                start_run_index = run_index
                start_offset = (
                    start_position - current_position
                )

            if end_position <= next_position:
                end_run_index = run_index
                end_offset = (
                    end_position - current_position
                )
                break

            current_position = next_position

        if (
            start_run_index is None
            or end_run_index is None
        ):
            continue

        start_run = runs[start_run_index]
        end_run = runs[end_run_index]
        prefix = start_run.text[:start_offset]
        suffix = end_run.text[end_offset:]

        start_run.text = (
            prefix + replacement + suffix
        )

        for run_index in range(
            start_run_index + 1,
            end_run_index + 1,
        ):
            runs[run_index].text = ""


def render_docx_template(
    *,
    template_path: Path,
    output_path: Path,
    values: Mapping[str, Any],
    required_variables: list[str],
) -> None:
    try:
        document = Document(template_path)
    except Exception as error:
        raise InvalidDocxTemplateError from error

    template_variables = get_template_variables(
        document
    )
    required = {
        normalize_required_variable(variable)
        for variable in required_variables
    }
    variables_to_validate = (
        template_variables | required
    )
    missing_variables = sorted(
        variable
        for variable in variables_to_validate
        if (
            variable not in values
            or is_missing_value(values[variable])
        )
    )

    if missing_variables:
        raise MissingTemplateVariablesError(
            missing_variables
        )

    text_values = {
        variable: value_to_text(values[variable])
        for variable in template_variables
    }

    for paragraph in iter_document_paragraphs(
        document
    ):
        replace_paragraph_placeholders(
            paragraph,
            text_values,
        )

    remaining_variables = get_template_variables(
        document
    )

    if remaining_variables:
        raise MissingTemplateVariablesError(
            sorted(remaining_variables)
        )

    try:
        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )
        document.save(output_path)
    except Exception as error:
        if output_path.exists():
            output_path.unlink()

        raise InvalidDocxTemplateError from error
